import io
import tempfile
import unittest
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

import app as reporting
from openpyxl import load_workbook, Workbook
from werkzeug.security import generate_password_hash


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path_patch = patch.object(reporting, "DB_PATH", Path(self.temp.name) / "test.sqlite3")
        self.path_patch.start()
        reporting.app.config["TESTING"] = True
        with reporting.app.app_context():
            reporting.init_db()
            with reporting.db() as connection:
                connection.executemany("INSERT INTO municipalities(id,name) VALUES (?,?)", [(1, "Город А"), (2, "Город Б")])
                connection.executemany("INSERT INTO schools(id,municipality_id,name) VALUES (?,?,?)", [(1, 1, "Школа 1"), (2, 2, "Школа 2")])
                connection.executemany("INSERT INTO users(username,password_hash,full_name,role,municipality_id) VALUES (?,?,?,?,?)", [
                    ("admin", generate_password_hash("secretpassword12"), "Администратор", "admin", None),
                    ("city_a", generate_password_hash("secretpassword12"), "Координатор А", "municipality", 1),
                    ("city_b", generate_password_hash("secretpassword12"), "Координатор Б", "municipality", 2),
                ])
        self.client = reporting.app.test_client()

    def tearDown(self):
        self.path_patch.stop()
        self.temp.cleanup()

    def csrf(self):
        with self.client.session_transaction() as session:
            return session["csrf"]

    def login(self, name):
        self.client.get("/login")
        response = self.client.post("/login", data={"csrf": self.csrf(), "username": name, "password": "secretpassword12"}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)

    def valid(self, date=None):
        return {"csrf": self.csrf(), "date": date or reporting.today().isoformat(), "staff_1": "1.5", "occupied_1": "0.5",
                "comment_1": "Одна ставка вакантна", "advisor_name_1_0": "Иванова И.И.", "status_1_0": "working",
                "employment_1_0": "parttime", "rate_1_0": "0.5", "month_1_0": "2026-09", "salary_1_0": "10000.50",
                "payment_org_1_0": "1", "full_month_1_0": "1", "reward_1_0": "5000", "insurance_1_0": "1500",
                "average_1_0": "0", "paid_1_0": "5000", "reason_1_0": ""}

    def test_report_scope_excel_and_revision(self):
        self.login("city_a")
        self.assertIn("Школа 1".encode(), self.client.get("/my/report").data)
        self.assertNotIn("Школа 2".encode(), self.client.get("/my/report").data)
        self.assertEqual(self.client.get("/admin").status_code, 403)
        self.assertEqual(self.client.get("/admin/export").status_code, 403)
        self.assertEqual(self.client.get("/admin/edit/2").status_code, 403)
        result = self.client.post("/my/report", data=self.valid(), follow_redirects=True)
        self.assertIn("Еженедельный отчет сохранен".encode(), result.data)
        bad = self.valid()
        bad["rate_1_0"] = "1"
        self.client.post("/my/report", data=bad)
        with reporting.app.app_context():
            self.assertEqual(reporting.db().execute("SELECT COUNT(*) FROM report_audit").fetchone()[0], 0)
            self.assertEqual(reporting.db().execute("SELECT occupied_rate FROM advisors").fetchone()[0], "0.5")
        amended = self.valid()
        amended["paid_1_0"] = "0"
        amended["reason_1_0"] = "Перенос выплаты"
        self.client.post("/my/report", data=amended)
        with reporting.app.app_context():
            self.assertEqual(reporting.db().execute("SELECT COUNT(*) FROM reports").fetchone()[0], 1)
            self.assertEqual(reporting.db().execute("SELECT COUNT(*) FROM report_audit").fetchone()[0], 1)
            self.assertEqual(reporting.db().execute("SELECT paid FROM advisors").fetchone()[0], 0)
        self.client.post("/logout", data={"csrf": self.csrf()})
        self.login("admin")
        self.assertIn("Не сдан".encode(), self.client.get("/admin").data)
        report_date = reporting.today().isoformat()
        self.assertIn("Иванова".encode(), self.client.get("/admin/report/1?date=" + report_date).data)
        exported = load_workbook(io.BytesIO(self.client.get("/admin/export?date=" + report_date).data))
        self.assertEqual(exported.sheetnames, ["Муниципалитеты", "Организации", "Советники и выплаты"])
        self.assertEqual(exported.worksheets[0][2][7].value, 1)
        self.assertEqual(exported.worksheets[2][2][2].value, "Иванова И.И.")

    def test_validation_and_access(self):
        self.assertEqual(self.client.post("/login", data={"username": "admin"}).status_code, 400)
        self.login("city_a")
        self.assertEqual(self.client.get("/my/report?date=2999-01-01").status_code, 400)
        self.assertEqual(self.client.get("/admin/directory").status_code, 403)
        for modification in ({"staff_1": "-1"}, {"staff_1": "2", "comment_1": ""}, {"reward_1_0": "4999"},
                             {"advisor_name_2_0": "Чужой"}, {"paid_1_0": "0"}):
            self.client.post("/my/report", data={**self.valid(), **modification})
        with reporting.app.app_context():
            self.assertEqual(reporting.db().execute("SELECT COUNT(*) FROM reports").fetchone()[0], 0)

    def test_period_and_excel_import(self):
        self.login("admin")
        book = Workbook()
        book.active.append(["№ п/п", "АТЕ", "Название образовательной организации", "Сокращенное название", "Тип",
                            "Осуществляет образовательный процесс/не осуществляет образовательный процесс/другое"])
        book.active.append([1, "Город В", "Колледж 3", "К3", "ПОО", "Осуществляет образовательный процесс"])
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post("/admin/directory", data={"csrf": self.csrf(), "action": "import", "file": (output, "registry.xlsx")}, follow_redirects=True)
        self.assertIn("Колледж 3".encode(), response.data)
        report_date = reporting.today().isoformat()
        due = (reporting.today() + timedelta(days=2)).isoformat()
        self.client.post("/admin/period", data={"csrf": self.csrf(), "report_date": report_date, "due_date": due})
        self.assertIn(due.encode(), self.client.get("/admin").data)
        with reporting.app.app_context():
            school = reporting.db().execute("SELECT kind,operation_status FROM schools WHERE name='Колледж 3'").fetchone()
            self.assertEqual(tuple(school), ("ПОО", "Осуществляет образовательный процесс"))


if __name__ == "__main__":
    unittest.main()
