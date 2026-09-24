"""Weekly municipality reporting for educational advisors."""

import io
import json
import os
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo

import click
from flask import Flask, abort, flash, g, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash

ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("REPORT_DB", ROOT / "reports.sqlite3"))
secret_file = ROOT / ".secret-key"
secret = os.environ.get("SECRET_KEY")
if not secret:
    if not secret_file.exists():
        secret_file.write_text(secrets.token_hex(32), encoding="ascii")
    secret = secret_file.read_text(encoding="ascii").strip()

app = Flask(__name__)
app.config.update(SECRET_KEY=secret, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=os.environ.get("HTTPS_ONLY") == "1",
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=12), MAX_CONTENT_LENGTH=5 * 1024 * 1024)


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_error):
    connection = g.pop("db", None)
    if connection is not None:
        connection.close()


def init_db():
    with db() as connection:
        connection.executescript("""
            CREATE TABLE IF NOT EXISTS municipalities (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE);
            CREATE TABLE IF NOT EXISTS schools (
                id INTEGER PRIMARY KEY, municipality_id INTEGER NOT NULL REFERENCES municipalities(id),
                name TEXT NOT NULL, UNIQUE(municipality_id, name));
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE, password_hash TEXT NOT NULL,
                full_name TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin', 'municipality')),
                municipality_id INTEGER REFERENCES municipalities(id));
            CREATE TABLE IF NOT EXISTS reports (
                id INTEGER PRIMARY KEY, municipality_id INTEGER NOT NULL REFERENCES municipalities(id),
                report_date TEXT NOT NULL, responsible_name TEXT NOT NULL, updated_at TEXT NOT NULL,
                UNIQUE(municipality_id, report_date));
            CREATE TABLE IF NOT EXISTS report_rows (
                report_id INTEGER NOT NULL REFERENCES reports(id) ON DELETE CASCADE,
                school_id INTEGER NOT NULL REFERENCES schools(id), assigned INTEGER NOT NULL DEFAULT 0,
                working INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(report_id, school_id));
            CREATE TABLE IF NOT EXISTS periods (
                report_date TEXT PRIMARY KEY, due_date TEXT NOT NULL, UNIQUE(due_date));
            CREATE TABLE IF NOT EXISTS advisors (
                id INTEGER PRIMARY KEY, report_id INTEGER NOT NULL REFERENCES reports(id) ON DELETE CASCADE,
                school_id INTEGER NOT NULL REFERENCES schools(id), full_name TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('working','leave','sick','maternity')),
                occupied_rate TEXT NOT NULL, employment TEXT NOT NULL,
                salary INTEGER NOT NULL, allowance INTEGER NOT NULL, bonus INTEGER NOT NULL,
                other INTEGER NOT NULL, paid INTEGER NOT NULL, payment_month TEXT NOT NULL,
                nonpayment_reason TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS report_audit (
                id INTEGER PRIMARY KEY, report_id INTEGER NOT NULL REFERENCES reports(id),
                changed_at TEXT NOT NULL, actor TEXT NOT NULL, snapshot TEXT NOT NULL);
            CREATE INDEX IF NOT EXISTS advisors_report ON advisors(report_id, school_id);
        """)
        columns = {row["name"] for row in connection.execute("PRAGMA table_info(report_rows)")}
        for name, definition in (("staff_units", "TEXT NOT NULL DEFAULT '0'"),
                                 ("occupied_rates", "TEXT NOT NULL DEFAULT '0'"),
                                 ("comment", "TEXT NOT NULL DEFAULT ''")):
            if name not in columns:
                connection.execute(f"ALTER TABLE report_rows ADD COLUMN {name} {definition}")
        if "kind" not in {row["name"] for row in connection.execute("PRAGMA table_info(schools)")}:
            connection.execute("ALTER TABLE schools ADD COLUMN kind TEXT NOT NULL DEFAULT 'Школа'")
        school_columns = {row["name"] for row in connection.execute("PRAGMA table_info(schools)")}
        if "operation_status" not in school_columns:
            connection.execute("ALTER TABLE schools ADD COLUMN operation_status TEXT NOT NULL DEFAULT ''")
        advisor_columns = {row["name"] for row in connection.execute("PRAGMA table_info(advisors)")}
        for name, definition in (("reward", "INTEGER NOT NULL DEFAULT 0"),
                                 ("insurance", "INTEGER NOT NULL DEFAULT 0"),
                                 ("average_earnings", "INTEGER NOT NULL DEFAULT 0"),
                                 ("payment_org", "INTEGER NOT NULL DEFAULT 0"),
                                 ("full_month", "INTEGER NOT NULL DEFAULT 0")):
            if name not in advisor_columns:
                connection.execute(f"ALTER TABLE advisors ADD COLUMN {name} {definition}")


@app.cli.command("init-db")
def init_db_command():
    init_db()
    click.echo("База создана.")


@app.cli.command("create-admin")
@click.option("--username", prompt="Логин администратора")
@click.option("--full-name", prompt="ФИО администратора")
@click.password_option(confirmation_prompt=True)
def create_admin(username, full_name, password):
    if len(password) < 12:
        raise click.ClickException("Пароль должен содержать не менее 12 символов.")
    init_db()
    try:
        with db() as connection:
            connection.execute("INSERT INTO users(username,password_hash,full_name,role) VALUES (?,?,?,'admin')",
                               (username.strip(), generate_password_hash(password), full_name.strip()))
    except sqlite3.IntegrityError as error:
        raise click.ClickException("Логин уже занят.") from error
    click.echo("Администратор создан.")


def import_directory_workbook(book, connection):
    sheet = book.active
    records = list(sheet.values)
    if not records:
        raise ValueError("В книге нет строк.")
    headers = [str(cell or "").strip() for cell in records[0]]
    if all(name in headers for name in ("АТЕ", "Название образовательной организации", "Тип")):
        municipality_index = headers.index("АТЕ")
        school_index = headers.index("Название образовательной организации")
        kind_index = headers.index("Тип")
        status_index = next((i for i, value in enumerate(headers) if value.startswith("Осуществляет образовательный процесс/")), None)
    elif tuple(value.casefold() for value in headers[:3]) == ("муниципалитет", "организация", "тип"):
        municipality_index, school_index, kind_index, status_index = 0, 1, 2, None
    else:
        raise ValueError("Не найдены столбцы АТЕ, Название образовательной организации и Тип.")
    if len(records) > 10001:
        raise ValueError("Не более 10000 строк в одном файле.")
    imported = 0
    for row_number, row in enumerate(records[1:], start=2):
        values = list(row) + [None] * len(headers)
        municipality = str(values[municipality_index] or "").strip()
        school = str(values[school_index] or "").strip()
        kind = str(values[kind_index] or "").strip()
        status = str(values[status_index] or "").strip() if status_index is not None else ""
        if not municipality and not school and not kind:
            continue
        if not municipality or not school or not kind or len(municipality) > 200 or len(school) > 500 or len(kind) > 80:
            raise ValueError(f"Проверьте обязательные поля в строке {row_number}.")
        connection.execute("INSERT OR IGNORE INTO municipalities(name) VALUES (?)", (municipality,))
        municipality_id = connection.execute("SELECT id FROM municipalities WHERE name=?", (municipality,)).fetchone()["id"]
        connection.execute("""INSERT INTO schools(municipality_id,name,kind,operation_status) VALUES (?,?,?,?)
            ON CONFLICT(municipality_id,name) DO UPDATE SET kind=excluded.kind,operation_status=excluded.operation_status""",
                           (municipality_id, school, kind, status))
        imported += 1
    return imported


@app.cli.command("import-directory")
@click.option("--path", type=click.Path(exists=True, dir_okay=False, path_type=Path),
              default=ROOT.parent / "советники" / "Актуальный_справочник_10_09_2026.xlsx", show_default=True)
def import_directory_command(path):
    init_db()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        with db() as connection:
            imported = import_directory_workbook(book, connection)
    finally:
        book.close()
    click.echo(f"Загружено организаций: {imported}.")


def today():
    return datetime.now(ZoneInfo("Europe/Moscow")).date()


def current_period():
    latest = db().execute("SELECT report_date FROM periods WHERE report_date <= ? ORDER BY report_date DESC LIMIT 1",
                          (today().isoformat(),)).fetchone()
    thursday = (today() - timedelta(days=(today().weekday() - 3) % 7)).isoformat()
    return max(thursday, latest["report_date"]) if latest else thursday


def selected_date():
    value = request.values.get("date") or current_period()
    try:
        parsed = date.fromisoformat(value)
    except (ValueError, TypeError):
        abort(400, "Неверная дата.")
    if parsed > today():
        abort(400, "Будущая отчетная дата недоступна.")
    return parsed.isoformat()


def due_date(report_date):
    period = db().execute("SELECT due_date FROM periods WHERE report_date = ?", (report_date,)).fetchone()
    return period["due_date"] if period else report_date


@app.before_request
def load_user_and_check_csrf():
    g.user = None
    if session.get("user_id"):
        g.user = db().execute("SELECT * FROM users WHERE id = ?", (session["user_id"],)).fetchone()
        if g.user is None:
            session.clear()
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    if request.method in ("POST", "PUT", "PATCH", "DELETE"):
        if not secrets.compare_digest(session["csrf"], request.form.get("csrf", "")):
            abort(400, "Недействительный токен формы.")


@app.context_processor
def template_values():
    return {"csrf": session.get("csrf"), "current_user": g.get("user"), "today": today().isoformat()}


def login_required(role=None):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if g.user is None:
                return redirect(url_for("login"))
            if role and g.user["role"] != role:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


@app.route("/")
def home():
    if g.user is None:
        return redirect(url_for("login"))
    return redirect(url_for("admin_dashboard" if g.user["role"] == "admin" else "municipal_dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = db().execute("SELECT * FROM users WHERE username = ?", (request.form.get("username", "").strip(),)).fetchone()
        if user and check_password_hash(user["password_hash"], request.form.get("password", "")):
            session.clear()
            session.update(user_id=user["id"], csrf=secrets.token_urlsafe(32))
            session.permanent = True
            return redirect(url_for("home"))
        flash("Неверный логин или пароль.", "error")
    return render_template("login.html")


@app.post("/logout")
@login_required()
def logout():
    session.clear()
    return redirect(url_for("login"))


def number(value, label, integer=False):
    try:
        result = Decimal(str(value).replace(",", "."))
        if not result.is_finite() or result < 0 or result > 100000 or result.as_tuple().exponent < -2:
            raise ValueError()
        if integer and result != result.to_integral_value():
            raise ValueError()
        return int(result) if integer else result
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"Некорректное значение: {label} (0–100000, не более двух знаков после запятой).") from None


def money(value, label):
    return int(number(value, label) * 100)


def parse_entries(schools):
    entries, advisors = [], []
    payment_names = set()
    municipality_id = schools[0]["municipality_id"] if schools else None
    school_ids = {s["id"] for s in schools}
    for key in request.form:
        match = re.fullmatch(r"advisor_name_(\d+)_(\d+)", key)
        if match and int(match.group(1)) not in school_ids:
            raise ValueError("Посторонняя организация в отчете.")
    for school in schools:
        sid = school["id"]
        staff = number(request.form.get(f"staff_{sid}"), f"штатные единицы: {school['name']}")
        occupied = number(request.form.get(f"occupied_{sid}"), f"занятые ставки: {school['name']}")
        comment = request.form.get(f"comment_{sid}", "").strip()
        if occupied > staff or len(comment) > 1000:
            raise ValueError(f"Занятые ставки не могут превышать штатные; комментарий до 1000 символов: {school['name']}.")
        indices = {int(m.group(1)) for key in request.form if (m := re.fullmatch(rf"advisor_name_{sid}_(\d+)", key))}
        if len(indices) > 1000:
            raise ValueError("Слишком много советников в одной организации.")
        names = set()
        rate_total = Decimal("0")
        for index in sorted(indices):
            prefix = f"{sid}_{index}"
            name = request.form.get(f"advisor_name_{prefix}", "").strip()
            if not name or len(name) > 200 or name.casefold() in names:
                raise ValueError(f"Укажите уникальное ФИО каждого советника: {school['name']}.")
            names.add(name.casefold())
            status = request.form.get(f"status_{prefix}")
            employment = request.form.get(f"employment_{prefix}")
            month = request.form.get(f"month_{prefix}", "")
            if status not in ("working", "leave", "sick", "maternity") or employment not in ("main", "parttime") or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
                raise ValueError(f"Проверьте статус, вид занятости и месяц выплат: {name}.")
            rate = number(request.form.get(f"rate_{prefix}"), f"ставка: {name}")
            if rate <= 0:
                raise ValueError(f"Ставка должна быть больше нуля: {name}.")
            rate_total += rate
            amounts = [money(request.form.get(f"{field}_{prefix}"), f"{field}: {name}")
                       for field in ("reward", "insurance", "average", "paid")]
            payment_org = request.form.get(f"payment_org_{prefix}") == "1"
            full_month = request.form.get(f"full_month_{prefix}") == "1"
            if full_month and status != "working":
                raise ValueError(f"Полностью отработанный месяц возможен только при статусе «Работает»: {name}.")
            if payment_org:
                if name.casefold() in payment_names:
                    raise ValueError(f"Для одного советника можно выбрать только одну организацию выплаты: {name}.")
                payment_names.add(name.casefold())
                duplicate = db().execute("""SELECT s.name FROM advisors a JOIN reports r ON r.id=a.report_id
                    JOIN schools s ON s.id=a.school_id WHERE lower(a.full_name)=lower(?) AND a.payment_month=?
                    AND a.payment_org=1 AND r.municipality_id<>? LIMIT 1""",
                    (name, month, municipality_id)).fetchone()
                if duplicate:
                    raise ValueError(f"Для {name} уже указана организация выплаты за {month}: {duplicate['name']}.")
                if full_month and amounts[0] != 500000:
                    raise ValueError(f"При полностью отработанном месяце вознаграждение должно составлять 5 000,00 руб.: {name}.")
            elif any(amounts):
                raise ValueError(f"Финансовые суммы разрешены только для выбранной организации выплаты: {name}.")
            reason = request.form.get(f"reason_{prefix}", "").strip()
            if len(reason) > 500 or (sum(amounts[:3]) > 0 and amounts[3] == 0 and not reason):
                raise ValueError(f"Укажите причину отсутствия выплаты (до 500 символов): {name}.")
            advisors.append((sid, name, status, str(rate), employment, *amounts, month, reason,
                             int(payment_org), int(full_month)))
        if rate_total != occupied:
            raise ValueError(f"Сумма ставок советников должна совпадать с занятыми ставками: {school['name']}.")
        if staff > occupied and not comment:
            raise ValueError(f"Поясните незанятые ставки в комментарии: {school['name']}.")
        entries.append((sid, str(staff), str(occupied), comment))
    return entries, advisors


def report_data(municipality_id, report_date):
    report = db().execute("SELECT * FROM reports WHERE municipality_id = ? AND report_date = ?",
                          (municipality_id, report_date)).fetchone()
    rows = {} if not report else {r["school_id"]: dict(r) for r in db().execute(
        "SELECT * FROM report_rows WHERE report_id = ?", (report["id"],))}
    advisors = {} if not report else {}
    if report:
        for advisor in db().execute("SELECT * FROM advisors WHERE report_id = ? ORDER BY school_id,id", (report["id"],)):
            advisors.setdefault(advisor["school_id"], []).append(dict(advisor))
    return report, rows, advisors


@app.route("/my")
@login_required("municipality")
def municipal_dashboard():
    municipality_id = g.user["municipality_id"]
    municipality = db().execute("SELECT * FROM municipalities WHERE id = ?", (municipality_id,)).fetchone()
    history = db().execute("""SELECT r.report_date,r.updated_at,COUNT(rr.school_id) organizations,
        COALESCE(SUM(CAST(rr.staff_units AS REAL)),0) staff,
        (SELECT COUNT(*) FROM advisors a WHERE a.report_id = r.id AND a.status = 'working') working
        FROM reports r LEFT JOIN report_rows rr ON rr.report_id = r.id
        WHERE r.municipality_id = ? GROUP BY r.id ORDER BY r.report_date DESC LIMIT 100""", (municipality_id,)).fetchall()
    period = current_period()
    return render_template("my.html", municipality=municipality, history=history, period=period,
                           due=due_date(period))


@app.route("/my/report", methods=["GET", "POST"])
@app.route("/admin/edit/<int:municipality_id>", methods=["GET", "POST"])
@login_required()
def municipal_report(municipality_id=None):
    if municipality_id is None:
        if g.user["role"] != "municipality":
            abort(403)
        municipality_id = g.user["municipality_id"]
    elif g.user["role"] != "admin":
        abort(403)
    municipality = db().execute("SELECT * FROM municipalities WHERE id = ?", (municipality_id,)).fetchone()
    if not municipality:
        abort(404)
    report_date = selected_date()
    schools = db().execute("SELECT * FROM schools WHERE municipality_id = ? ORDER BY name", (municipality_id,)).fetchall()
    report, rows, advisors = report_data(municipality_id, report_date)
    if request.method == "POST":
        try:
            if not schools:
                raise ValueError("Сначала добавьте организации в справочник.")
            entries, people = parse_entries(schools)
            previous = db().execute("SELECT report_date FROM reports WHERE municipality_id = ? AND report_date < ? ORDER BY report_date DESC LIMIT 1",
                                    (municipality_id, report_date)).fetchone()
            if previous:
                _, old_rows, _ = report_data(municipality_id, previous["report_date"])
                changed = any(str(old_rows.get(sid, {}).get("staff_units", "0")) != staff for sid, staff, _, _ in entries)
                if changed and not request.form.get("change_reason", "").strip():
                    raise ValueError("При изменении штатных единиц укажите причину изменения относительно прошлого отчета.")
            now = datetime.now(ZoneInfo("Europe/Moscow")).isoformat(timespec="seconds")
            with db() as connection:
                if report:
                    snapshot = {"rows": list(rows.values()), "advisors": [a for group in advisors.values() for a in group],
                                "change_reason": request.form.get("change_reason", "").strip()}
                    connection.execute("INSERT INTO report_audit(report_id,changed_at,actor,snapshot) VALUES (?,?,?,?)",
                                       (report["id"], now, g.user["full_name"], json.dumps(snapshot, ensure_ascii=False)))
                    connection.execute("UPDATE reports SET responsible_name=?,updated_at=? WHERE id=?",
                                       (g.user["full_name"], now, report["id"]))
                    report_id = report["id"]
                    connection.execute("DELETE FROM advisors WHERE report_id=?", (report_id,))
                    connection.execute("DELETE FROM report_rows WHERE report_id=?", (report_id,))
                else:
                    report_id = connection.execute("INSERT INTO reports(municipality_id,report_date,responsible_name,updated_at) VALUES (?,?,?,?)",
                                                   (municipality_id, report_date, g.user["full_name"], now)).lastrowid
                connection.executemany("INSERT INTO report_rows(report_id,school_id,assigned,working,staff_units,occupied_rates,comment) VALUES (?,?,0,0,?,?,?)",
                                       ((report_id, *entry) for entry in entries))
                connection.executemany("""INSERT INTO advisors(report_id,school_id,full_name,status,occupied_rate,employment,
                    reward,insurance,average_earnings,paid,payment_month,nonpayment_reason,payment_org,full_month,
                    salary,allowance,bonus,other) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,0,0,0)""",
                                       ((report_id, *person) for person in people))
            flash("Еженедельный отчет сохранен.", "success")
            return redirect(url_for("admin_report", municipality_id=municipality_id, date=report_date)
                            if g.user["role"] == "admin" else url_for("municipal_dashboard"))
        except ValueError as error:
            flash(str(error), "error")
    return render_template("report.html", date=report_date, due=due_date(report_date), schools=schools,
                           municipality=municipality, rows=rows, advisors=advisors, report=report,
                           month=report_date[:7])


def summary_for(report_date):
    return db().execute("""SELECT m.id,m.name,r.id report_id,r.responsible_name,r.updated_at,
        COUNT(rr.school_id) school_count,COALESCE(SUM(CAST(rr.staff_units AS REAL)),0) staff,
        COALESCE(SUM(CAST(rr.occupied_rates AS REAL)),0) occupied,
        (SELECT COUNT(*) FROM advisors a WHERE a.report_id=r.id AND a.status='working') working,
        (SELECT COALESCE(SUM(a.reward+a.average_earnings),0) FROM advisors a WHERE a.report_id=r.id) accrued,
        (SELECT COALESCE(SUM(a.insurance),0) FROM advisors a WHERE a.report_id=r.id) insurance,
        (SELECT COALESCE(SUM(a.paid),0) FROM advisors a WHERE a.report_id=r.id) paid
        FROM municipalities m LEFT JOIN reports r ON r.municipality_id=m.id AND r.report_date=?
        LEFT JOIN report_rows rr ON rr.report_id=r.id GROUP BY m.id ORDER BY m.name""", (report_date,)).fetchall()


@app.route("/admin")
@login_required("admin")
def admin_dashboard():
    report_date = selected_date()
    summary = summary_for(report_date)
    periods = db().execute("SELECT * FROM periods ORDER BY report_date DESC LIMIT 30").fetchall()
    return render_template("admin.html", date=report_date, due=due_date(report_date), summary=summary, periods=periods)


@app.post("/admin/period")
@login_required("admin")
def set_period():
    try:
        report_date = date.fromisoformat(request.form.get("report_date", ""))
        due = date.fromisoformat(request.form.get("due_date", ""))
        if due < report_date or report_date > today() + timedelta(days=365):
            raise ValueError()
        with db() as connection:
            connection.execute("INSERT INTO periods(report_date,due_date) VALUES (?,?) ON CONFLICT(report_date) DO UPDATE SET due_date=excluded.due_date",
                               (report_date.isoformat(), due.isoformat()))
        flash("Отчетная дата и срок сдачи сохранены.", "success")
    except (ValueError, sqlite3.IntegrityError):
        flash("Проверьте даты: срок сдачи не раньше отчетной даты.", "error")
    return redirect(url_for("admin_dashboard"))


@app.route("/admin/report/<int:municipality_id>")
@login_required("admin")
def admin_report(municipality_id):
    report_date = selected_date()
    municipality = db().execute("SELECT * FROM municipalities WHERE id=?", (municipality_id,)).fetchone()
    if not municipality:
        abort(404)
    report, rows, advisors = report_data(municipality_id, report_date)
    schools = db().execute("SELECT * FROM schools WHERE municipality_id=? ORDER BY name", (municipality_id,)).fetchall()
    audit = [] if not report else [dict(entry, snapshot=json.loads(entry["snapshot"])) for entry in db().execute(
        "SELECT changed_at,actor,snapshot FROM report_audit WHERE report_id=? ORDER BY id DESC", (report["id"],))]
    return render_template("admin_report.html", date=report_date, municipality=municipality, report=report,
                           rows=rows, advisors=advisors, schools=schools, audit=audit)


def safe_cell(value):
    return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value


@app.route("/admin/export")
@login_required("admin")
def export():
    report_date = selected_date()
    book = Workbook()
    overview = book.active
    overview.title = "Муниципалитеты"
    overview.append(["Отчетная дата", "Срок сдачи", "Муниципалитет", "Статус", "Организаций",
                     "Штатных единиц", "Занятых ставок", "Работающих (чел.)", "Вакансий (ставок)",
                     "Укомплектованность, %", "Вознаграждение и средний заработок, руб.", "Страховые отчисления, руб.",
                     "Выплачено советникам, руб.", "Не выплачено, руб."])
    summary = summary_for(report_date)
    for r in summary:
        overview.append([report_date, due_date(report_date), safe_cell(r["name"]), "Сдан" if r["report_id"] else "Не сдан",
                         r["school_count"] if r["report_id"] else None, r["staff"] if r["report_id"] else None,
                         r["occupied"] if r["report_id"] else None, r["working"] if r["report_id"] else None,
                         round(r["staff"] - r["occupied"], 2) if r["report_id"] else None,
                         round(r["occupied"] / r["staff"] * 100, 2) if r["report_id"] and r["staff"] else None,
                         r["accrued"] / 100 if r["report_id"] else None, r["insurance"] / 100 if r["report_id"] else None,
                         r["paid"] / 100 if r["report_id"] else None,
                         (r["accrued"] - r["paid"]) / 100 if r["report_id"] else None])
    submitted = [r for r in summary if r["report_id"]]
    overview.append([report_date, "", "ИТОГО ПО РЕГИОНУ", f"{len(submitted)}/{len(summary)}", sum(r["school_count"] for r in submitted),
                     sum(r["staff"] for r in submitted), sum(r["occupied"] for r in submitted), sum(r["working"] for r in submitted),
                     round(sum(r["staff"] - r["occupied"] for r in submitted), 2), None,
                     sum(r["accrued"] for r in submitted) / 100, sum(r["insurance"] for r in submitted) / 100,
                     sum(r["paid"] for r in submitted) / 100,
                     sum(r["accrued"] - r["paid"] for r in submitted) / 100])
    org = book.create_sheet("Организации")
    org.append(["Муниципалитет", "Организация", "Тип", "Штатных единиц", "Занятых ставок", "Советников (чел.)",
                "Работающих (чел.)", "Вакансий", "Комментарий", "Вознаграждение и средний заработок, руб.",
                "Страховые отчисления, руб.", "Выплачено советникам, руб."])
    detail = book.create_sheet("Советники и выплаты")
    detail.append(["Муниципалитет", "Организация", "ФИО", "Статус", "Занятость", "Ставка", "Месяц выплат",
                   "Организация выплаты", "Полностью отработан месяц", "Вознаграждение", "Страховые отчисления",
                   "Средний заработок", "Начислено советнику", "Выплачено советнику", "Не выплачено", "Причина"])
    rows = db().execute("""SELECT m.name municipality,s.name school,s.kind,rr.*,r.id report_id FROM reports r
        JOIN municipalities m ON m.id=r.municipality_id JOIN report_rows rr ON rr.report_id=r.id
        JOIN schools s ON s.id=rr.school_id WHERE r.report_date=? ORDER BY m.name,s.name""", (report_date,)).fetchall()
    for row in rows:
        people = db().execute("SELECT * FROM advisors WHERE report_id=? AND school_id=? ORDER BY full_name",
                              (row["report_id"], row["school_id"])).fetchall()
        accrued = sum(p["reward"] + p["average_earnings"] for p in people)
        insurance = sum(p["insurance"] for p in people)
        paid = sum(p["paid"] for p in people)
        org.append([safe_cell(row["municipality"]), safe_cell(row["school"]), safe_cell(row["kind"]),
                    float(row["staff_units"]), float(row["occupied_rates"]), len(people),
                    sum(p["status"] == "working" for p in people),
                    round(float(row["staff_units"]) - float(row["occupied_rates"]), 2), safe_cell(row["comment"]),
                    accrued / 100, insurance / 100, paid / 100])
        for p in people:
            detail.append([safe_cell(row["municipality"]), safe_cell(row["school"]), safe_cell(p["full_name"]),
                           p["status"], p["employment"], float(p["occupied_rate"]), p["payment_month"],
                           "Да" if p["payment_org"] else "Нет", "Да" if p["full_month"] else "Нет",
                           p["reward"] / 100, p["insurance"] / 100, p["average_earnings"] / 100,
                           accrued_p := (p["reward"] + p["average_earnings"]) / 100,
                           p["paid"] / 100, accrued_p - p["paid"] / 100, safe_cell(p["nonpayment_reason"])])
    from openpyxl.styles import Font, PatternFill
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174D5A")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(55, max(15, max(len(str(c.value or "")) for c in column) + 2))
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"advisors-{report_date}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/admin/directory", methods=["GET", "POST"])
@login_required("admin")
def directory():
    if request.method == "POST":
        action = request.form.get("action")
        name = request.form.get("name", "").strip()
        try:
            with db() as connection:
                if action in ("municipality", "school", "rename_municipality", "rename_school"):
                    if not name or len(name) > 250:
                        raise ValueError("Укажите название (до 250 символов).")
                    if action == "municipality":
                        connection.execute("INSERT INTO municipalities(name) VALUES (?)", (name,))
                    elif action == "school":
                        kind = request.form.get("kind", "").strip()
                        if not kind or len(kind) > 80:
                            raise ValueError("Укажите тип организации.")
                        connection.execute("INSERT INTO schools(municipality_id,name,kind) VALUES (?,?,?)",
                                           (int(request.form.get("municipality_id", "")), name, kind))
                    else:
                        table = "municipalities" if action == "rename_municipality" else "schools"
                        if not connection.execute(f"UPDATE {table} SET name=? WHERE id=?",
                                                  (name, int(request.form.get("id", "")))).rowcount:
                            raise ValueError("Запись не найдена.")
                elif action == "user":
                    username = request.form.get("username", "").strip()
                    full_name = request.form.get("full_name", "").strip()
                    password = request.form.get("password", "")
                    if not username or len(username) > 100 or not full_name or len(full_name) > 200 or len(password) < 12:
                        raise ValueError("Укажите логин, ФИО и пароль не короче 12 символов.")
                    connection.execute("INSERT INTO users(username,password_hash,full_name,role,municipality_id) VALUES (?,?,?,'municipality',?)",
                                       (username, generate_password_hash(password), full_name, int(request.form.get("municipality_id", ""))))
                elif action == "import":
                    upload = request.files.get("file")
                    if not upload or not upload.filename.lower().endswith(".xlsx"):
                        raise ValueError("Загрузите файл .xlsx.")
                    book = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
                    try:
                        import_directory_workbook(book, connection)
                    finally:
                        book.close()
                else:
                    abort(400)
            flash("Справочник обновлен.", "success")
            return redirect(url_for("directory"))
        except (ValueError, sqlite3.IntegrityError) as error:
            flash(f"Не удалось обновить справочник: {error}", "error")
    municipalities = db().execute("SELECT * FROM municipalities ORDER BY name").fetchall()
    schools = db().execute("SELECT s.*,m.name municipality FROM schools s JOIN municipalities m ON m.id=s.municipality_id ORDER BY m.name,s.name").fetchall()
    users = db().execute("SELECT u.username,u.full_name,m.name municipality FROM users u JOIN municipalities m ON m.id=u.municipality_id ORDER BY m.name,u.username").fetchall()
    return render_template("directory.html", municipalities=municipalities, schools=schools, users=users)


if __name__ == "__main__":
    app.run(debug=False)
